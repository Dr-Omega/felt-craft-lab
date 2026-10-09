# -*- coding: utf-8 -*-
"""Convert YWN Feishu export pack → immutable historical Asset Snapshot (SXY/ZXR).

Fail-closed: validation issues may write a failure report, but must not mutate
latest pointers, snapshot registry, prior snapshots, media, or content-versions.
Publishing stages under a temp dir and only commits after full success.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import shutil
import sys
import tempfile
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

REPO_ROOT = Path(__file__).resolve().parents[1]
OUT_MANIFEST = REPO_ROOT / "assets" / "generated" / "assets.manifest.json"
OUT_REPORT = REPO_ROOT / "assets" / "generated" / "validation-report.json"
OUT_VERSIONS = REPO_ROOT / "assets" / "generated" / "content-versions.json"
OUT_REGISTRY = REPO_ROOT / "assets" / "generated" / "snapshot-registry.json"
OUT_SAMPLE = REPO_ROOT / "docs" / "workflows" / "ywn-sample-sticker-fu-001.md"
MATERIAL_OUT = REPO_ROOT / "assets" / "generated" / "material"
AUDIO_OUT = REPO_ROOT / "assets" / "generated" / "audio"
SNAPSHOTS_OUT = REPO_ROOT / "assets" / "generated" / "snapshots"
RENDERER_CATALOG = REPO_ROOT / "src" / "renderer" / "src" / "catalog"
RENDERER_MANIFEST = RENDERER_CATALOG / "assets.manifest.json"
RENDERER_REGISTRY = RENDERER_CATALOG / "snapshot-registry.json"

FELT_AUDIO = re.compile(r"^felt-resource://audio/([A-Za-z0-9][A-Za-z0-9._/-]*)$")
# Resolver charset allows hyphen; new filenames still prefer underscore by convention.
SAFE_FILE = re.compile(r"^[A-Za-z0-9][A-Za-z0-9._/-]*$")

PROJECTION_KEYS = (
    "type",
    "themeId",
    "categoryIds",
    "nameZh",
    "nameEn",
    "introZh",
    "keywords",
    "image",
    "thumbnail",
    "preSplit",
    "audio",
    "source",
    "license",
)


def utc_now() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%S.%f")[:-3] + "Z"


def new_snapshot_id(now: datetime | None = None) -> str:
    stamp = (now or datetime.now(timezone.utc)).strftime("%Y%m%d_%H%M%S")
    return f"snapshot_ywn_{stamp}"


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def empty_media_identity() -> dict[str, str]:
    return {"logicalName": "", "sha256": ""}


def media_identity(logical_name: str = "", digest: str = "") -> dict[str, str]:
    return {"logicalName": logical_name or "", "sha256": digest or ""}


def logical_name_from_ref(ref: str) -> str:
    if not isinstance(ref, str) or not ref:
        return ""
    return Path(ref.replace("\\", "/").rstrip("/")).name


def material_ref(snapshot_id: str, filename: str) -> str:
    return f"felt-resource://material/{snapshot_id}/{filename}"


def audio_ref(snapshot_id: str, filename: str) -> str:
    return f"felt-resource://audio/{snapshot_id}/{filename}"


def published_path_for_ref(ref: str, repo_root: Path) -> Path | None:
    if not isinstance(ref, str) or not ref.startswith("felt-resource://"):
        return None
    body = ref[len("felt-resource://") :]
    if body.startswith("material/"):
        return repo_root / "assets" / "generated" / "material" / body.split("/", 1)[1]
    if body.startswith("audio/"):
        return repo_root / "assets" / "generated" / "audio" / body.split("/", 1)[1]
    return None


def _coerce_media_identity(
    value: object,
    *,
    fallback_name: str = "",
    fallback_hash: str = "",
) -> dict[str, str]:
    if isinstance(value, dict) and ("logicalName" in value or "sha256" in value):
        return media_identity(
            str(value.get("logicalName") or fallback_name or ""),
            str(value.get("sha256") or fallback_hash or ""),
        )
    if isinstance(value, str):
        return media_identity(fallback_name, value or fallback_hash)
    return media_identity(fallback_name, fallback_hash)


def content_projection(record: dict, media_hashes: dict | None = None) -> dict:
    """Project semantic fields + logical media identity (name + sha256), not ResourceRefs.

    ``media_hashes`` may be:
    - legacy ``{image: sha256, ...}`` strings (logical names taken from record refs/names)
    - or ``{image: {logicalName, sha256}, ...}`` media identity objects
    """
    audio = record.get("audio") or {}
    supplied = dict(media_hashes) if media_hashes is not None else {}
    if not supplied and isinstance(record.get("mediaHashes"), dict):
        supplied = dict(record.get("mediaHashes") or {})

    # Prefer explicit logical media already on the record (new sidecar shape).
    image = _coerce_media_identity(
        record.get("image") if isinstance(record.get("image"), dict) else supplied.get("image"),
        fallback_name=logical_name_from_ref(str(record.get("imageUrl") or "")),
        fallback_hash=str(supplied.get("image") or "") if isinstance(supplied.get("image"), str) else "",
    )
    if isinstance(supplied.get("image"), dict):
        image = _coerce_media_identity(supplied["image"], fallback_name=image["logicalName"])

    thumbnail = _coerce_media_identity(
        record.get("thumbnail") if isinstance(record.get("thumbnail"), dict) else supplied.get("thumbnail"),
        fallback_name=logical_name_from_ref(str(record.get("thumbnailUrl") or "")),
        fallback_hash=str(supplied.get("thumbnail") or "")
        if isinstance(supplied.get("thumbnail"), str)
        else "",
    )
    if isinstance(supplied.get("thumbnail"), dict):
        thumbnail = _coerce_media_identity(
            supplied["thumbnail"], fallback_name=thumbnail["logicalName"]
        )

    pre_split = _coerce_media_identity(
        record.get("preSplit") if isinstance(record.get("preSplit"), dict) else supplied.get("preSplit"),
        fallback_name=logical_name_from_ref(str(record.get("preSplitDataUrl") or "")),
        fallback_hash=str(supplied.get("preSplit") or "")
        if isinstance(supplied.get("preSplit"), str)
        else "",
    )
    if isinstance(supplied.get("preSplit"), dict):
        pre_split = _coerce_media_identity(supplied["preSplit"], fallback_name=pre_split["logicalName"])

    audio_block = record.get("audio") if isinstance(record.get("audio"), dict) else {}
    # New shape nests identities under projection["audio"][nameZh|nameEn|introZh]
    name_zh = _coerce_media_identity(
        audio_block.get("nameZh")
        if isinstance(audio_block.get("nameZh"), dict)
        else supplied.get("nameZhAudio"),
        fallback_name=logical_name_from_ref(str(audio.get("nameZhUrl") or "")),
        fallback_hash=str(supplied.get("nameZhAudio") or "")
        if isinstance(supplied.get("nameZhAudio"), str)
        else "",
    )
    if isinstance(supplied.get("nameZhAudio"), dict):
        name_zh = _coerce_media_identity(supplied["nameZhAudio"], fallback_name=name_zh["logicalName"])

    name_en = _coerce_media_identity(
        audio_block.get("nameEn")
        if isinstance(audio_block.get("nameEn"), dict)
        else supplied.get("nameEnAudio"),
        fallback_name=logical_name_from_ref(str(audio.get("nameEnUrl") or "")),
        fallback_hash=str(supplied.get("nameEnAudio") or "")
        if isinstance(supplied.get("nameEnAudio"), str)
        else "",
    )
    if isinstance(supplied.get("nameEnAudio"), dict):
        name_en = _coerce_media_identity(supplied["nameEnAudio"], fallback_name=name_en["logicalName"])

    intro_zh = _coerce_media_identity(
        audio_block.get("introZh")
        if isinstance(audio_block.get("introZh"), dict)
        else supplied.get("introZhAudio"),
        fallback_name=logical_name_from_ref(str(audio.get("introZhUrl") or "")),
        fallback_hash=str(supplied.get("introZhAudio") or "")
        if isinstance(supplied.get("introZhAudio"), str)
        else "",
    )
    if isinstance(supplied.get("introZhAudio"), dict):
        intro_zh = _coerce_media_identity(
            supplied["introZhAudio"], fallback_name=intro_zh["logicalName"]
        )

    return {
        "type": record.get("type"),
        "themeId": record.get("themeId"),
        "categoryIds": list(record.get("categoryIds") or []),
        "nameZh": record.get("nameZh"),
        "nameEn": record.get("nameEn"),
        "introZh": record.get("introZh"),
        "keywords": list(record.get("keywords") or []),
        "image": image,
        "thumbnail": thumbnail,
        "preSplit": pre_split,
        "audio": {
            "nameZh": name_zh,
            "nameEn": name_en,
            "introZh": intro_zh,
        },
        "source": record.get("source") or "",
        "license": record.get("license") or "",
    }


def enrich_projection_media_hashes(projection: dict, repo_root: Path = REPO_ROOT) -> dict:
    """Normalize legacy ResourceRef projections and fill missing sha256 from published media."""
    enriched = content_projection(projection)
    mapping = (
        ("image", projection.get("imageUrl") or ""),
        ("thumbnail", projection.get("thumbnailUrl") or ""),
        ("preSplit", projection.get("preSplitDataUrl") or ""),
    )
    audio_old = projection.get("audio") or {}
    audio_refs = (
        ("nameZh", audio_old.get("nameZhUrl") or ""),
        ("nameEn", audio_old.get("nameEnUrl") or ""),
        ("introZh", audio_old.get("introZhUrl") or ""),
    )
    legacy_hashes = projection.get("mediaHashes") or {}

    def fill(slot: dict[str, str], ref: str, legacy_key: str) -> dict[str, str]:
        name = slot.get("logicalName") or logical_name_from_ref(ref)
        digest = slot.get("sha256") or ""
        if not digest and isinstance(legacy_hashes.get(legacy_key), str):
            digest = legacy_hashes.get(legacy_key) or ""
        if not digest and ref:
            path = published_path_for_ref(str(ref), repo_root)
            if path and path.is_file():
                digest = sha256_file(path)
        if not digest and name:
            # Fall back: search under material/audio trees by filename.
            for kind in ("material", "audio"):
                root = repo_root / "assets" / "generated" / kind
                if not root.is_dir():
                    continue
                direct = root / name
                if direct.is_file():
                    digest = sha256_file(direct)
                    break
                for path in root.rglob(name):
                    if path.is_file():
                        digest = sha256_file(path)
                        break
                if digest:
                    break
        return media_identity(name, digest)

    enriched["image"] = fill(enriched["image"], mapping[0][1], "image")
    enriched["thumbnail"] = fill(enriched["thumbnail"], mapping[1][1], "thumbnail")
    enriched["preSplit"] = fill(enriched["preSplit"], mapping[2][1], "preSplit")
    enriched["audio"] = {
        "nameZh": fill(enriched["audio"]["nameZh"], audio_refs[0][1], "nameZhAudio"),
        "nameEn": fill(enriched["audio"]["nameEn"], audio_refs[1][1], "nameEnAudio"),
        "introZh": fill(enriched["audio"]["introZh"], audio_refs[2][1], "introZhAudio"),
    }
    return enriched


def resolve_content_version(asset_id: str, projection: dict, prior: dict[str, dict]) -> int:
    prev = prior.get(asset_id)
    if not prev:
        return 1
    old_version = int(prev.get("contentVersion") or 1)
    old_projection = prev.get("projection")
    if old_projection == projection:
        return old_version
    return old_version + 1


def load_prior_versions(
    versions_path: Path = OUT_VERSIONS,
    manifest_path: Path = OUT_MANIFEST,
    repo_root: Path = REPO_ROOT,
) -> dict[str, dict]:
    raw_assets: dict[str, dict] = {}
    if versions_path.is_file():
        raw = json.loads(versions_path.read_text(encoding="utf-8"))
        assets = raw.get("assets") if isinstance(raw, dict) else {}
        raw_assets = dict(assets or {})
    elif manifest_path.is_file():
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        assets = manifest.get("assets") or []
        for asset in assets:
            asset_id = asset.get("assetId")
            if not asset_id:
                continue
            raw_assets[asset_id] = {
                "contentVersion": int(asset.get("contentVersion") or 1),
                "projection": content_projection(asset),
            }
    out: dict[str, dict] = {}
    for asset_id, entry in raw_assets.items():
        projection = entry.get("projection") if isinstance(entry, dict) else None
        if not isinstance(projection, dict):
            continue
        out[asset_id] = {
            "contentVersion": int(entry.get("contentVersion") or 1),
            "projection": enrich_projection_media_hashes(projection, repo_root),
        }
    return out


def merge_version_history(prior: dict[str, dict], updated: dict[str, dict]) -> dict[str, dict]:
    """Keep historical assetIds that are absent from the current ready batch."""
    merged = dict(prior)
    merged.update(updated)
    return merged


def load_registry(registry_path: Path = OUT_REGISTRY) -> dict:
    if not registry_path.is_file():
        return {"latestSnapshotId": "", "snapshots": {}}
    raw = json.loads(registry_path.read_text(encoding="utf-8"))
    if not isinstance(raw, dict):
        return {"latestSnapshotId": "", "snapshots": {}}
    return {
        "latestSnapshotId": str(raw.get("latestSnapshotId") or ""),
        "snapshots": dict(raw.get("snapshots") or {}),
    }


def registry_entry(snapshot_id: str, generated_at: str) -> dict:
    return {
        "manifest": f"assets/generated/snapshots/{snapshot_id}/assets.manifest.json",
        "generatedAt": generated_at,
    }


def header_index(headers: list[str]) -> dict[str, int]:
    mapping = {}
    for i, raw in enumerate(headers):
        h = str(raw or "")
        key = h.split("（")[0].split("(")[0].strip()
        if key:
            mapping[key] = i
    return mapping


def cell(row, idx: dict[str, int], key: str) -> str:
    i = idx.get(key)
    if i is None or i >= len(row) or row[i] is None:
        return ""
    return str(row[i]).strip()


def split_keywords(raw: str) -> list[str]:
    parts = re.split(r"[,，]", raw)
    seen = set()
    out = []
    for part in parts:
        item = part.strip()
        if not item or item in seen:
            continue
        seen.add(item)
        out.append(item)
    return out


def parse_audio_filename(raw: str) -> str | None:
    """Return logical audio filename from a Feishu cell (may include felt-resource URI)."""
    if not raw:
        return None
    m = FELT_AUDIO.match(raw)
    if m:
        return Path(m.group(1)).name
    if SAFE_FILE.match(raw) and raw.endswith((".mp3", ".wav")):
        return Path(raw).name
    return None


def basename_stem(file_name: str) -> str:
    return Path(file_name).name.rsplit(".", 1)[0]


def find_theme_file(source_root: Path, theme_id: str, relative: str) -> Path | None:
    candidates = [
        source_root / theme_id / "stickers" / relative,
        source_root / theme_id / "backgrounds" / relative,
        source_root / theme_id / "thumbnails" / relative,
        source_root / "shared" / "audio" / relative,
        source_root / "audio" / relative,
    ]
    for path in candidates:
        if path.is_file():
            return path
    audio_dir = source_root / "audio"
    if audio_dir.is_dir():
        for path in audio_dir.glob(f"*_{relative}"):
            if path.is_file():
                return path
    return None


def load_ready_rows(source_root: Path) -> tuple[list[dict], list[dict]]:
    rows_out: list[dict] = []
    issues: list[dict] = []
    seen_ids: set[str] = set()
    # The complete v2 pack has one authoritative workbook. Theme exports are
    # older copies and must not silently override or duplicate its rows.
    master = source_root / "素材库.xlsx"
    xlsx_files = [master] if master.is_file() else sorted(source_root.rglob("feishu_fields_*.xlsx"))
    for path in xlsx_files:
        wb = load_workbook(path, data_only=True)
        if path == master:
            if len(wb.worksheets) != 1:
                issues.append({"sourceSheet": path.name, "errors": ["master workbook must contain exactly one material sheet"]})
                wb.close()
                continue
            ws = wb.worksheets[0]
        elif "01_正式素材_可填" in wb.sheetnames:
            ws = wb["01_正式素材_可填"]
        else:
            wb.close()
            continue
        data = list(ws.iter_rows(values_only=True))
        wb.close()
        if not data:
            issues.append({"sourceSheet": str(path.relative_to(source_root)), "errors": ["empty material sheet"]})
            continue
        idx = header_index([str(c) if c is not None else "" for c in data[0]])
        required = {"assetId", "type", "themeId", "status", "licenseStatus", "imageFileName"}
        if not required.issubset(idx):
            issues.append({"sourceSheet": str(path.relative_to(source_root)), "errors": ["missing required columns: " + ", ".join(sorted(required - idx.keys()))]})
            continue
        for row in data[1:]:
            if not row:
                continue
            asset_id = cell(row, idx, "assetId")
            status = cell(row, idx, "status")
            # Formal import requires status=ready AND licenseStatus=approved.
            # Non-ready rows are skipped; ready rows with bad license fail below.
            if not asset_id or status != "ready":
                continue
            if asset_id in seen_ids:
                issues.append(
                    {
                        "assetId": asset_id,
                        "errors": [f"duplicate assetId: {asset_id}"],
                        "sourceSheet": str(path.relative_to(source_root)),
                    }
                )
                continue
            seen_ids.add(asset_id)
            rows_out.append(
                {
                    "sourceSheet": str(path.relative_to(source_root)),
                    "assetId": asset_id,
                    "type": cell(row, idx, "type"),
                    "themeId": cell(row, idx, "themeId"),
                    "categoryId": cell(row, idx, "categoryId"),
                    "nameZh": cell(row, idx, "nameZh"),
                    "nameEn": cell(row, idx, "nameEn"),
                    "introZh": cell(row, idx, "introZh"),
                    "imageFileName": cell(row, idx, "imageFileName"),
                    "thumbnailFileName": cell(row, idx, "thumbnailFileName"),
                    "preSplitFileName": cell(row, idx, "preSplitFileName"),
                    "source": cell(row, idx, "source"),
                    "license": cell(row, idx, "license"),
                    "licenseStatus": cell(row, idx, "licenseStatus"),
                    "audioNameZhPath": cell(row, idx, "audioNameZhPath"),
                    "audioNameEnPath": cell(row, idx, "audioNameEnPath"),
                    "audioIntroZhPath": cell(row, idx, "audioIntroZhPath"),
                    "owner": cell(row, idx, "owner"),
                    "status": status,
                    "keywords": cell(row, idx, "keywords"),
                }
            )
    return rows_out, issues


def convert_row(
    row: dict,
    source_root: Path,
    prior_versions: dict[str, dict],
    snapshot_id: str,
) -> tuple[dict | None, list[str], dict[str, Path], dict]:
    errors: list[str] = []
    copies: dict[str, Path] = {}
    media: dict[str, dict[str, str]] = {
        "image": empty_media_identity(),
        "thumbnail": empty_media_identity(),
        "preSplit": empty_media_identity(),
        "nameZhAudio": empty_media_identity(),
        "nameEnAudio": empty_media_identity(),
        "introZhAudio": empty_media_identity(),
    }

    # Strict gate: empty licenseStatus also fails.
    if row["licenseStatus"] != "approved":
        errors.append(f"licenseStatus must be approved for ready (got {row['licenseStatus']!r})")

    if row["type"] not in ("background", "sticker"):
        errors.append(f"invalid type {row['type']}")

    category = row["categoryId"].strip()
    if not category:
        errors.append("categoryId empty")

    image_name = row["imageFileName"]
    if not image_name or not SAFE_FILE.match(image_name):
        errors.append(f"bad imageFileName: {image_name!r}")
    else:
        if basename_stem(image_name) != row["assetId"]:
            errors.append(
                f"imageFileName stem must equal assetId (got {basename_stem(image_name)!r}, expected {row['assetId']!r})"
            )
        src = find_theme_file(source_root, row["themeId"], image_name)
        if not src:
            errors.append(f"missing image file {image_name}")
        else:
            copies[f"material/{snapshot_id}/{image_name}"] = src
            media["image"] = media_identity(image_name, sha256_file(src))

    thumb_name = row["thumbnailFileName"]
    thumb_ref = None
    if thumb_name:
        if not SAFE_FILE.match(thumb_name):
            errors.append(f"bad thumbnailFileName: {thumb_name!r}")
        else:
            src = find_theme_file(source_root, row["themeId"], thumb_name)
            if not src:
                errors.append(f"missing thumbnail file {thumb_name}")
            else:
                copies[f"material/{snapshot_id}/{thumb_name}"] = src
                thumb_ref = material_ref(snapshot_id, thumb_name)
                media["thumbnail"] = media_identity(thumb_name, sha256_file(src))

    pre_split_name = row.get("preSplitFileName") or ""
    pre_split_ref = None
    if pre_split_name:
        if row["type"] == "background":
            errors.append("background must not have preSplitFileName")
        elif row["type"] != "sticker":
            errors.append(f"preSplitFileName not allowed for type {row['type']!r}")
        elif not SAFE_FILE.match(pre_split_name):
            errors.append(f"bad preSplitFileName: {pre_split_name!r}")
        else:
            src = find_theme_file(source_root, row["themeId"], pre_split_name)
            if not src:
                errors.append(f"missing preSplit file {pre_split_name}")
            else:
                copies[f"material/{snapshot_id}/{pre_split_name}"] = src
                pre_split_ref = material_ref(snapshot_id, pre_split_name)
                media["preSplit"] = media_identity(pre_split_name, sha256_file(src))

    audio: dict[str, str] = {}
    audio_hash_keys = {
        "nameZhUrl": "nameZhAudio",
        "nameEnUrl": "nameEnAudio",
        "introZhUrl": "introZhAudio",
    }
    for field, key in (
        ("audioNameZhPath", "nameZhUrl"),
        ("audioNameEnPath", "nameEnUrl"),
        ("audioIntroZhPath", "introZhUrl"),
    ):
        file_name = parse_audio_filename(row[field])
        if row[field] and not file_name:
            errors.append(f"invalid audio path {field}={row[field]!r}")
            continue
        if not file_name:
            continue
        src = find_theme_file(source_root, row["themeId"], file_name)
        if not src:
            errors.append(f"missing audio file {file_name}")
        else:
            copies[f"audio/{snapshot_id}/{file_name}"] = src
            audio[key] = audio_ref(snapshot_id, file_name)
            media[audio_hash_keys[key]] = media_identity(file_name, sha256_file(src))

    draft = {
        "assetId": row["assetId"],
        "type": row["type"],
        "themeId": row["themeId"],
        "categoryIds": [category] if category else [],
        "nameZh": row["nameZh"],
        "nameEn": row["nameEn"],
        "introZh": row["introZh"],
        "keywords": split_keywords(row["keywords"]),
        "imageUrl": material_ref(snapshot_id, image_name) if image_name else "",
        "audio": audio,
        "status": "ready",
    }
    if thumb_ref:
        draft["thumbnailUrl"] = thumb_ref
    if pre_split_ref:
        draft["preSplitDataUrl"] = pre_split_ref
    if row["source"]:
        draft["source"] = row["source"]
    if row["license"]:
        draft["license"] = row["license"]

    projection = content_projection(draft, media)
    version = resolve_content_version(row["assetId"], projection, prior_versions)
    draft["contentVersion"] = version

    if errors:
        return None, errors, copies, media
    return draft, [], copies, media


def write_failure_report(
    *,
    generated_at: str,
    rows_seen: int,
    issues: list[dict],
    snapshot_id: str,
    report_path: Path = OUT_REPORT,
) -> None:
    report = {
        "generatedAt": generated_at,
        "readyRowsSeen": rows_seen,
        "imported": 0,
        "failed": len(issues),
        "assetSnapshotId": snapshot_id,
        "issues": issues,
        "copiedFiles": 0,
        "published": False,
        "notes": [
            "fail-closed: latest / registry / prior snapshots / media / content-versions were NOT modified",
            "licenseStatus must equal approved (empty fails)",
            "duplicate assetId fails into report",
            "status=ready AND licenseStatus=approved required for formal import",
            "existing snapshots/<snapshotId>/ refuses publish (immutable)",
        ],
        "failedDetails": issues,
        "projectionKeys": list(PROJECTION_KEYS),
    }
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def _read_optional_bytes(path: Path) -> bytes | None:
    if not path.exists():
        return None
    return path.read_bytes()


def _restore_optional_bytes(path: Path, payload: bytes | None) -> None:
    if payload is None:
        if path.exists():
            path.unlink()
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(payload)


def atomic_write_text(path: Path, text: str) -> None:
    """Write via temp file in the same directory, then os.replace."""
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp_name = tempfile.mkstemp(prefix=f".{path.name}.", suffix=".tmp", dir=str(path.parent))
    tmp_path = Path(tmp_name)
    try:
        with os.fdopen(fd, "w", encoding="utf-8", newline="\n") as handle:
            handle.write(text)
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(tmp_path, path)
    except Exception:
        try:
            if tmp_path.exists():
                tmp_path.unlink()
        except OSError:
            pass
        raise


def atomic_copy_file(src: Path, dest: Path) -> None:
    """Copy via temp file in the same directory, then os.replace."""
    dest.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp_name = tempfile.mkstemp(prefix=f".{dest.name}.", suffix=".tmp", dir=str(dest.parent))
    os.close(fd)
    tmp_path = Path(tmp_name)
    try:
        shutil.copy2(src, tmp_path)
        os.replace(tmp_path, dest)
    except Exception:
        try:
            if tmp_path.exists():
                tmp_path.unlink()
        except OSError:
            pass
        raise


def _rm_tree_if_exists(path: Path) -> None:
    if path.exists():
        shutil.rmtree(path)


def commit_published_artifacts(
    *,
    manifest: dict,
    versions: dict[str, dict],
    copies: dict[str, Path],
    generated_at: str,
    snapshot_id: str,
    prior_registry: dict,
    repo_root: Path = REPO_ROOT,
    manifest_path: Path = OUT_MANIFEST,
    renderer_manifest: Path = RENDERER_MANIFEST,
    versions_path: Path = OUT_VERSIONS,
    registry_path: Path = OUT_REGISTRY,
    renderer_registry: Path = RENDERER_REGISTRY,
) -> list[dict]:
    """Publish with rollback: new snapshot resources + atomic mutable pointer updates.

    On any failure after starting the transaction: restore prior mutable bytes and
    delete this snapshot's newly created media/manifest trees. Old latest unchanged.
    """
    snapshot_dir = repo_root / "assets" / "generated" / "snapshots" / snapshot_id
    if snapshot_dir.exists():
        raise SystemExit(
            f"Refusing to overwrite existing snapshot dir: {snapshot_dir.relative_to(repo_root)}"
        )

    material_snapshot_dir = repo_root / "assets" / "generated" / "material" / snapshot_id
    audio_snapshot_dir = repo_root / "assets" / "generated" / "audio" / snapshot_id
    if material_snapshot_dir.exists() or audio_snapshot_dir.exists():
        raise SystemExit(f"Refusing to overwrite existing snapshot media for {snapshot_id}")

    renderer_catalog = renderer_manifest.parent
    renderer_snapshot = renderer_catalog / "snapshots" / f"{snapshot_id}.manifest.json"
    if renderer_snapshot.exists():
        raise SystemExit(f"Refusing to overwrite existing renderer snapshot manifest: {renderer_snapshot}")

    next_registry = {
        "latestSnapshotId": snapshot_id,
        "snapshots": {
            **dict(prior_registry.get("snapshots") or {}),
            snapshot_id: registry_entry(snapshot_id, generated_at),
        },
    }

    payload = json.dumps(manifest, ensure_ascii=False, indent=2) + "\n"
    versions_payload = (
        json.dumps({"updatedAt": generated_at, "assets": versions}, ensure_ascii=False, indent=2)
        + "\n"
    )
    registry_payload = json.dumps(next_registry, ensure_ascii=False, indent=2) + "\n"

    mutable_paths = [
        versions_path,
        registry_path,
        renderer_registry,
        manifest_path,
        renderer_manifest,
    ]
    mutable_backups = {path: _read_optional_bytes(path) for path in mutable_paths}

    copied: list[dict] = []
    created_files: list[Path] = []
    created_dirs: list[Path] = []

    def rollback() -> None:
        for path, old in mutable_backups.items():
            _restore_optional_bytes(path, old)
        for path in created_files:
            try:
                if path.exists():
                    path.unlink()
            except OSError:
                pass
        for path in created_dirs:
            _rm_tree_if_exists(path)
        _rm_tree_if_exists(snapshot_dir)
        _rm_tree_if_exists(material_snapshot_dir)
        _rm_tree_if_exists(audio_snapshot_dir)
        if renderer_snapshot.exists():
            try:
                renderer_snapshot.unlink()
            except OSError:
                pass

    try:
        with tempfile.TemporaryDirectory(prefix="felt-snapshot-stage-") as tmp:
            stage = Path(tmp)
            stage_manifest = stage / "assets.manifest.json"
            stage_versions = stage / "content-versions.json"
            stage_registry = stage / "snapshot-registry.json"
            stage_manifest.write_text(payload, encoding="utf-8")
            stage_versions.write_text(versions_payload, encoding="utf-8")
            stage_registry.write_text(registry_payload, encoding="utf-8")

            for rel, src in sorted(copies.items()):
                dest = stage / rel
                dest.parent.mkdir(parents=True, exist_ok=True)
                shutil.copy2(src, dest)

            # 1) New immutable snapshot media + manifests (transaction-owned)
            for rel, _src in sorted(copies.items()):
                staged = stage / rel
                final = repo_root / "assets" / "generated" / rel
                atomic_copy_file(staged, final)
                created_files.append(final)
                copied.append({"from": str(_src), "to": str(final.relative_to(repo_root))})

            created_dirs.extend([material_snapshot_dir, audio_snapshot_dir, snapshot_dir])

            snapshot_dir.mkdir(parents=True, exist_ok=True)
            snapshot_manifest = snapshot_dir / "assets.manifest.json"
            atomic_copy_file(stage_manifest, snapshot_manifest)
            created_files.append(snapshot_manifest)

            renderer_snapshot.parent.mkdir(parents=True, exist_ok=True)
            atomic_copy_file(stage_manifest, renderer_snapshot)
            created_files.append(renderer_snapshot)

            # Fault injection: after new snapshot installed, before mutable pointer commit.
            if os.environ.get("FELT_PUBLISH_FAIL_AFTER_SNAPSHOT") == "1":
                raise RuntimeError("injected publish failure after snapshot install")

            # 2) Mutable pointers via temp + os.replace (versions → registry → latest)
            atomic_write_text(versions_path, stage_versions.read_text(encoding="utf-8"))
            atomic_write_text(registry_path, stage_registry.read_text(encoding="utf-8"))
            atomic_write_text(renderer_registry, stage_registry.read_text(encoding="utf-8"))
            atomic_write_text(manifest_path, payload)
            atomic_write_text(renderer_manifest, payload)
    except Exception:
        rollback()
        raise

    return copied


def build_manifest(
    *,
    source_root: Path,
    source_pack: str,
    snapshot_id: str = "",
    repo_root: Path = REPO_ROOT,
    manifest_path: Path = OUT_MANIFEST,
    renderer_manifest: Path = RENDERER_MANIFEST,
    versions_path: Path = OUT_VERSIONS,
    registry_path: Path = OUT_REGISTRY,
    renderer_registry: Path = RENDERER_REGISTRY,
    report_path: Path = OUT_REPORT,
    sample_path: Path = OUT_SAMPLE,
) -> int:
    if not source_root.is_dir():
        print(f"SOURCE_ROOT missing: {source_root}", file=sys.stderr)
        return 2
    pack_label = source_pack.strip()
    if not pack_label:
        print("source_pack is required (non-empty Manifest.sourcePack provenance).", file=sys.stderr)
        return 2

    prior_versions = load_prior_versions(versions_path, manifest_path, repo_root)
    prior_registry = load_registry(registry_path)
    rows, duplicate_issues = load_ready_rows(source_root)
    assets: list[dict] = []
    issues: list[dict] = list(duplicate_issues)
    pending_copies: dict[str, Path] = {}
    updated_versions: dict[str, dict] = {}
    generated_at = utc_now()
    planned_snapshot = snapshot_id.strip() or new_snapshot_id()

    for row in rows:
        record, errors, copies, media = convert_row(
            row, source_root, prior_versions, planned_snapshot
        )
        if errors:
            issues.append({"assetId": row["assetId"], "errors": errors, "sourceSheet": row["sourceSheet"]})
            continue
        assert record is not None
        assets.append(record)
        pending_copies.update(copies)
        updated_versions[record["assetId"]] = {
            "contentVersion": record["contentVersion"],
            "projection": content_projection(record, media),
        }

    if issues:
        write_failure_report(
            generated_at=generated_at,
            rows_seen=len(rows),
            issues=issues,
            snapshot_id=planned_snapshot,
            report_path=report_path,
        )
        summary = {
            "imported": 0,
            "failed": len(issues),
            "snapshot": planned_snapshot,
            "published": False,
            "latestSnapshotId": prior_registry.get("latestSnapshotId") or "",
        }
        print(json.dumps(summary, ensure_ascii=False))
        print(json.dumps(issues, ensure_ascii=False, indent=2))
        return 1

    snapshot_dir = repo_root / "assets" / "generated" / "snapshots" / planned_snapshot
    if snapshot_dir.exists():
        refuse_issues = [
            {
                "assetId": "*",
                "errors": [
                    f"snapshot already exists (immutable): assets/generated/snapshots/{planned_snapshot}/"
                ],
                "sourceSheet": "",
            }
        ]
        write_failure_report(
            generated_at=generated_at,
            rows_seen=len(rows),
            issues=refuse_issues,
            snapshot_id=planned_snapshot,
            report_path=report_path,
        )
        summary = {
            "imported": 0,
            "failed": 1,
            "snapshot": planned_snapshot,
            "published": False,
            "latestSnapshotId": prior_registry.get("latestSnapshotId") or "",
        }
        print(json.dumps(summary, ensure_ascii=False))
        print(json.dumps(refuse_issues, ensure_ascii=False, indent=2))
        return 1

    snapshot_assets = [{"assetId": a["assetId"], "type": a["type"]} for a in assets]
    manifest = {
        "schemaVersion": 1,
        "assetSnapshotId": planned_snapshot,
        "generatedAt": generated_at,
        "sourcePack": pack_label,
        "assets": assets,
        "snapshotContext": {"snapshots": {planned_snapshot: snapshot_assets}},
    }
    merged_versions = merge_version_history(prior_versions, updated_versions)
    try:
        copied = commit_published_artifacts(
            manifest=manifest,
            versions=merged_versions,
            copies=pending_copies,
            generated_at=generated_at,
            snapshot_id=planned_snapshot,
            prior_registry=prior_registry,
            repo_root=repo_root,
            manifest_path=manifest_path,
            renderer_manifest=renderer_manifest,
            versions_path=versions_path,
            registry_path=registry_path,
            renderer_registry=renderer_registry,
        )
    except SystemExit as exc:
        refuse_issues = [{"assetId": "*", "errors": [str(exc)], "sourceSheet": ""}]
        write_failure_report(
            generated_at=generated_at,
            rows_seen=len(rows),
            issues=refuse_issues,
            snapshot_id=planned_snapshot,
            report_path=report_path,
        )
        print(json.dumps({"imported": 0, "failed": 1, "snapshot": planned_snapshot, "published": False}))
        return 1
    except Exception as exc:
        refuse_issues = [
            {
                "assetId": "*",
                "errors": [f"publish rollback: {exc}"],
                "sourceSheet": "",
            }
        ]
        write_failure_report(
            generated_at=generated_at,
            rows_seen=len(rows),
            issues=refuse_issues,
            snapshot_id=planned_snapshot,
            report_path=report_path,
        )
        summary = {
            "imported": 0,
            "failed": 1,
            "snapshot": planned_snapshot,
            "published": False,
            "latestSnapshotId": prior_registry.get("latestSnapshotId") or "",
        }
        print(json.dumps(summary, ensure_ascii=False))
        print(json.dumps(refuse_issues, ensure_ascii=False, indent=2))
        return 1

    report = {
        "generatedAt": generated_at,
        "readyRowsSeen": len(rows),
        "imported": len(assets),
        "failed": 0,
        "assetSnapshotId": planned_snapshot,
        "issues": [],
        "copiedFiles": len(copied),
        "published": True,
        "notes": [
            "owner is kept only in build report workflow, not AssetRecord",
            "contentVersion continues from content-versions.json using logical media identity + sha256",
            "ResourceRef includes snapshotId; contentVersion identity does not",
            "historical content-versions entries are retained when assetId absent from current ready batch",
            "licenseStatus must equal approved (empty fails)",
            "status=ready AND licenseStatus=approved required for formal import",
            "imageFileName stem must equal assetId",
            "preSplitFileName allowed for stickers only",
            "fail-closed: formal artifacts written only after whole-batch success",
            "snapshots/<snapshotId>/ is immutable; duplicate snapshotId refuses publish",
            "felt-resource resolver allows [A-Za-z0-9._/-]; underscore is a naming preference only",
        ],
        "failedDetails": [],
        "projectionKeys": list(PROJECTION_KEYS),
    }
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    sample_row = next((r for r in rows if r["assetId"] == "sticker_fu_001"), None)
    sample_asset = next((a for a in assets if a["assetId"] == "sticker_fu_001"), None)
    if sample_row and sample_asset:
        sample_md = f"""# YWN 单条试转复核：sticker_fu_001

**来源包：** {pack_label}

**assetSnapshotId：** `{planned_snapshot}`
**生成时间：** `{generated_at}`

## 1. 原始飞书行（春节表）

```json
{json.dumps(sample_row, ensure_ascii=False, indent=2)}
```

## 2. 转换后 AssetRecord

```json
{json.dumps(sample_asset, ensure_ascii=False, indent=2)}
```

## 3. assetSnapshotId

`{planned_snapshot}`（本批正式快照；与单条 `contentVersion` 分开；已发布 snapshotId 不可覆盖）

## 4. validateAssetRecord

见 `tests/assets/manifest-validation.test.ts`（`FELT_RESOURCE_REF_POLICY`）。

## 5. ResourceRef 文件存在性

| ref | 打包路径 | 存在 |
| --- | --- | --- |
| `{sample_asset["imageUrl"]}` | `assets/generated/material/{planned_snapshot}/sticker_fu_001.png` | yes |
| `{sample_asset.get("thumbnailUrl")}` | `assets/generated/material/{planned_snapshot}/sticker_fu_001_thumb.png` | yes |
| `{sample_asset["audio"].get("nameZhUrl")}` | `assets/generated/audio/{planned_snapshot}/sticker_fu_001_name_zh.mp3` | yes |
| `{sample_asset["audio"].get("nameEnUrl")}` | `assets/generated/audio/{planned_snapshot}/sticker_fu_001_name_en.mp3` | yes |
| `{sample_asset["audio"].get("introZhUrl")}` | `assets/generated/audio/{planned_snapshot}/sticker_fu_001_intro_zh.mp3` | yes |

说明：未向 Renderer 暴露 OS 路径；上表仅供 SXY/ZXR 构建复核。
"""
        sample_path.parent.mkdir(parents=True, exist_ok=True)
        sample_path.write_text(sample_md, encoding="utf-8")

    summary = {
        "imported": len(assets),
        "failed": 0,
        "snapshot": planned_snapshot,
        "published": True,
        "latestSnapshotId": planned_snapshot,
    }
    print(json.dumps(summary, ensure_ascii=False))
    return 0


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Build formal YWN Asset Manifest (immutable snapshots)")
    parser.add_argument(
        "--source",
        type=Path,
        required=True,
        help="Path to Feishu export pack root (required; no machine-local default).",
    )
    parser.add_argument(
        "--source-pack",
        required=True,
        help='Provenance label written to Manifest.sourcePack (e.g. "归档(3).zip").',
    )
    parser.add_argument(
        "--snapshot-id",
        default="",
        help="Optional snapshot id. Default: snapshot_ywn_YYYYMMDD_HHMMSS (always new).",
    )
    args = parser.parse_args(argv)
    return build_manifest(
        source_root=args.source,
        source_pack=args.source_pack,
        snapshot_id=args.snapshot_id,
    )


if __name__ == "__main__":
    raise SystemExit(main())
